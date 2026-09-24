"""
Knowledge Base Module - Xử lý tài liệu và tìm kiếm vector (RAG)

Luồng hoạt động:
1. Giảng viên/Sinh viên upload tài liệu (PDF, DOCX, TXT, code)
2. Hệ thống cắt tài liệu thành các đoạn nhỏ (chunks)
3. Mỗi chunk được chuyển thành vector embedding bằng Gemini
4. Lưu vào ChromaDB (vector database chạy local, không cần server)
5. Khi sinh viên hỏi → tìm các chunk liên quan nhất → đưa vào prompt cho AI
"""

import os
import hashlib
from pathlib import Path
from typing import List

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from dotenv import load_dotenv

import chromadb
from chromadb.config import Settings

load_dotenv()

# Thư mục lưu trữ
UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

CHROMA_DIR = Path("chroma_db")
CHROMA_DIR.mkdir(exist_ok=True)


class KnowledgeBase:
    def __init__(self):
        api_key = os.getenv("GEMINI_API_KEY")

        # Embedding model - chuyển text thành vector
        self.embeddings = GoogleGenerativeAIEmbeddings(
            model="models/gemini-embedding-2",
            google_api_key=api_key,
            max_retries=1
        )

        # ChromaDB client - lưu trữ vector local
        self.chroma_client = chromadb.PersistentClient(path=str(CHROMA_DIR))
        self.collection = self.chroma_client.get_or_create_collection(
            name="winforms_knowledge",
            metadata={"description": "Tài liệu môn WinForms"}
        )

        # Text splitter - cắt tài liệu thành chunks
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            separators=["\n\n", "\n", ". ", " ", ""]
        )

    def _extract_text(self, file_path: str, filename: str) -> str:
        """Đọc nội dung file dựa trên loại file."""
        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text

        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(file_path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            return text

        elif ext in (".docx",):
            from docx import Document
            doc = Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs])

        elif ext in (".txt", ".cs", ".java", ".py", ".cpp", ".c", ".h", ".json", ".xml", ".md"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        else:
            # Thử đọc như text
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            except Exception:
                return ""

    async def add_document(self, file_path: str, filename: str) -> dict:
        """
        Xử lý một tài liệu: đọc → cắt chunks → embedding → lưu ChromaDB.
        Returns: {"status": ..., "chunks": số chunks đã lưu, "filename": ...}
        """
        # 1. Đọc nội dung
        text = self._extract_text(file_path, filename)
        if not text.strip():
            return {"status": "error", "message": "Không đọc được nội dung file", "filename": filename}

        # 2. Cắt thành chunks
        chunks = self.text_splitter.split_text(text)
        if not chunks:
            return {"status": "error", "message": "File không có đủ nội dung", "filename": filename}

        # 3. Tạo embedding và lưu vào ChromaDB
        doc_id = hashlib.md5(filename.encode()).hexdigest()[:8]
        ids = []
        documents = []
        metadatas = []

        for i, chunk in enumerate(chunks):
            chunk_id = f"{doc_id}_chunk_{i}"
            ids.append(chunk_id)
            documents.append(chunk)
            metadatas.append({
                "source": filename,
                "chunk_index": i,
                "total_chunks": len(chunks)
            })

        # Tạo embeddings
        embeddings = self.embeddings.embed_documents(documents)

        # Lưu vào ChromaDB (upsert để tránh duplicate)
        self.collection.upsert(
            ids=ids,
            documents=documents,
            embeddings=embeddings,
            metadatas=metadatas
        )

        return {
            "status": "success",
            "filename": filename,
            "chunks": len(chunks),
            "message": f"Đã xử lý '{filename}' thành {len(chunks)} đoạn kiến thức"
        }

    async def search(self, query: str, n_results: int = 5) -> List[dict]:
        """
        Tìm kiếm các đoạn kiến thức liên quan nhất với câu hỏi.
        Returns: List of {"text": ..., "source": ..., "score": ...}
        """
        if self.collection.count() == 0:
            return []

        # Tạo embedding cho câu hỏi
        query_embedding = self.embeddings.embed_query(query)

        # Tìm kiếm trong ChromaDB
        results = self.collection.query(
            query_embeddings=[query_embedding],
            n_results=min(n_results, self.collection.count())
        )

        # Format kết quả
        search_results = []
        if results and results["documents"]:
            for i, doc in enumerate(results["documents"][0]):
                search_results.append({
                    "text": doc,
                    "source": results["metadatas"][0][i].get("source", "Không rõ nguồn"),
                    "score": results["distances"][0][i] if results["distances"] else 0
                })

        return search_results

    def get_stats(self) -> dict:
        """Thống kê knowledge base."""
        count = self.collection.count()
        # Lấy danh sách các file đã upload
        sources = set()
        if count > 0:
            all_data = self.collection.get(include=["metadatas"])
            for meta in all_data["metadatas"]:
                sources.add(meta.get("source", "unknown"))

        return {
            "total_chunks": count,
            "total_documents": len(sources),
            "documents": list(sources)
        }


# Singleton instance
knowledge_base = KnowledgeBase()
